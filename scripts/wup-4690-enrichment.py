#!/usr/bin/env python3
"""Fail-closed enrichment of the exact 4,690 F21 extension rows.

F21 defines membership only. Population is sourced from the separately released
JRC GHS-WUP-MTUC R2025A statistics package. A match is accepted only when
country + normalized city name is unique and both datasets provide coordinates
within 1 degree. Missing/ambiguous/geographically inconsistent rows are never
imputed and cause the job to fail closed.
"""
import argparse,csv,hashlib,io,json,math,re,sys,zipfile
from pathlib import Path
import requests
from openpyxl import load_workbook
F21_URL='https://population.un.org/wup/assets/Download/Cities/WUP2025-F21-DEGURBA-Cities_Pop.xlsx'
JRC_URL='https://jeodpp.jrc.ec.europa.eu/ftp/jrc-opendata/GHSL/GHS_WUP_MTUC_GLOBE_R2025A/V1-1/GHS_WUP_MTUC_GLOBE_R2025A_V1_1_statistics.zip'
EXPECTED=4690; MAX_COORD_DEGREES=1.0

def norm(v):
 s='' if v is None else str(v); s=s.casefold().strip(); s=re.sub(r'[\u2018\u2019\u201a\u201b`]','\'',s); s=re.sub(r'[^\w\s\-\'&]',' ',s,flags=re.UNICODE); return ' '.join(s.split())
def get(url):
 r=requests.get(url,timeout=300); r.raise_for_status(); return r.content
def pick(headers,names):
 low={str(h).strip().casefold():h for h in headers}
 for n in names:
  if n.casefold() in low:return low[n.casefold()]
 return None
def coord(v):
 try:
  x=float(v); return x if math.isfinite(x) else None
 except:return None
def f21_records(raw):
 wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True); ws=wb['Data']; it=ws.iter_rows(values_only=True); header=[str(x).strip() if x is not None else '' for x in next(it)]
 ix={k:header.index(k) for k in ['Location','City_Name','City_Code','ISO3_Code','2025']}; lat=pick(header,['Latitude','lat','LAT']); lon=pick(header,['Longitude','lon','LON'])
 if not lat or not lon: raise SystemExit('f21-geography-columns-required')
 out=[]
 for idx,row in enumerate(it,start=1):
  if row[ix['2025']] not in (None,''): continue
  city=str(row[ix['City_Name']] or '').strip(); country=str(row[ix['Location']] or '').strip()
  if not city or not country: continue
  out.append({'country':country,'city':city,'wup_city_code':str(row[ix['City_Code']]).strip(),'iso3':str(row[ix['ISO3_Code']] or '').strip(),'source_row_index':idx,'lat':coord(row[header.index(lat)]),'lon':coord(row[header.index(lon)])})
 if len(out)!=EXPECTED: raise SystemExit(f'f21-extension-count:{len(out)}:expected:{EXPECTED}')
 if any(r['lat'] is None or r['lon'] is None for r in out): raise SystemExit('f21-extension-missing-geography')
 return out
def jrc_rows(raw):
 z=zipfile.ZipFile(io.BytesIO(raw)); names=[n for n in z.namelist() if n.lower().endswith(('.xlsx','.csv'))]
 if not names: raise SystemExit('jrc-statistics-no-machine-readable-table')
 names.sort(key=lambda n:(0 if 'MT_GLOBE' in n and n.lower().endswith('.xlsx') else 1,len(n)))
 with z.open(names[0]) as f:data=f.read()
 if names[0].lower().endswith('.csv'): return list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
 wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True); ws=wb.active; it=ws.iter_rows(values_only=True); header=[str(x).strip() if x is not None else '' for x in next(it)]; return [dict(zip(header,row)) for row in it]
def enrich(f21,jrc):
 keys=jrc[0].keys() if jrc else []; jcity=pick(keys,['UCname','city','name']); jcountry=pick(keys,['UNLocName','country']); jyear=pick(keys,['Year','year']); jpop=pick(keys,['POP','Population','population']); jlat=pick(keys,['Lat','lat','Latitude']); jlon=pick(keys,['Lon','lon','Longitude'])
 if not all([jcity,jcountry,jyear,jpop,jlat,jlon]): raise SystemExit(f'jrc-schema-mismatch:city={jcity}:country={jcountry}:year={jyear}:pop={jpop}:lat={jlat}:lon={jlon}')
 idx={}
 for r in jrc:
  try:year=int(float(r.get(jyear))); pop=float(r.get(jpop)); lat=coord(r.get(jlat)); lon=coord(r.get(jlon)))
  except:continue
  if year!=2025 or not math.isfinite(pop) or pop<=0 or pop>=50000 or lat is None or lon is None:continue
  idx.setdefault((norm(r.get(jcountry)),norm(r.get(jcity))),[]).append((r,pop,lat,lon))
 accepted=[]; unmatched=[]; ambiguous=[]; geo=[]
 for src in f21:
  cand=idx.get((norm(src['country']),norm(src['city'])),[])
  if len(cand)==0:unmatched.append(src);continue
  if len(cand)>1:ambiguous.append({'source':src,'candidates':len(cand)});continue
  r,pop,lat,lon=cand[0]; d=max(abs(src['lat']-lat),abs(src['lon']-lon))
  if d>MAX_COORD_DEGREES:geo.append({'source':src,'jrc_lat':lat,'jrc_lon':lon,'max_abs_coord_delta':d});continue
  accepted.append({**src,'population':int(round(pop)),'population_source':'JRC GHS-WUP-MTUC R2025A','population_source_year':2025,'jrc_city_name':r[jcity],'jrc_country':r[jcountry],'jrc_lat':lat,'jrc_lon':lon,'jrc_id':r.get('ID_UC_G0') or r.get('ID_MTUC'),'coordinate_max_abs_delta':d})
 report={'expected':EXPECTED,'accepted':len(accepted),'unmatched':len(unmatched),'ambiguous':len(ambiguous),'geographic_mismatches':len(geo),'status':'pass' if len(accepted)==EXPECTED and not unmatched and not ambiguous and not geo else 'fail'}
 return accepted,unmatched,ambiguous,geo,report
def main():
 ap=argparse.ArgumentParser();ap.add_argument('--out',required=True);ap.add_argument('--report',required=True);a=ap.parse_args(); f21=get(F21_URL);jrc=get(JRC_URL);f21r=f21_records(f21);jr=jrc_rows(jrc);accepted,unmatched,ambiguous,geo,report=enrich(f21r,jr);report.update({'f21_sha256':hashlib.sha256(f21).hexdigest(),'jrc_sha256':hashlib.sha256(jrc).hexdigest(),'f21_url':F21_URL,'jrc_url':JRC_URL,'max_coordinate_delta_degrees':MAX_COORD_DEGREES})
 if report['status']!='pass':
  Path(a.out).unlink(missing_ok=True);Path(a.report).write_text(json.dumps({**report,'unmatched_rows':unmatched[:100],'ambiguous_rows':ambiguous[:100],'geographic_mismatches':geo[:100]},indent=2)+'\n');raise SystemExit('4690-population-enrichment-failed-closed')
 accepted.sort(key=lambda r:r['source_row_index']);Path(a.out).write_text(json.dumps(accepted,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');Path(a.report).write_text(json.dumps(report,indent=2)+'\n');print(json.dumps(report,indent=2))
if __name__=='__main__':main()
