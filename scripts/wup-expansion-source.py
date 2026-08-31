#!/usr/bin/env python3
import hashlib
import json
import math
from pathlib import Path
from openpyxl import load_workbook

SOURCE_URL = 'https://population.un.org/wup/assets/Download/Cities/WUP2025-F21-DEGURBA-Cities_Pop.xlsx'
EXPECTED_TOTAL = 16828
EXPECTED_EXISTING = 12138
EXPECTED_EXTENSION = 4690
EXPECTED_SHA256 = '3a96030d87aec6c1c50f658d5321067d6345e1ab936c5d2854524f972caa75c0'


def build(source, output):
    raw = Path(source).read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != EXPECTED_SHA256:
        raise SystemExit(f'wup-f21-sha256-mismatch:{digest}:expected:{EXPECTED_SHA256}')
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb['Data']
    rows = list(ws.iter_rows(values_only=True))
    header = [str(x).strip() if x is not None else '' for x in rows[0]]
    required = {'Location','City_Name','City_Code','ISO3_Code','2025'}
    if not required.issubset(header):
        raise SystemExit(f'wup-f21-schema-mismatch:{sorted(required-set(header))}')
    ix = {name: header.index(name) for name in required}
    records=[]; existing=0; blank=0
    for source_row_index,row in enumerate(rows[1:], start=1):
        country=str(row[ix['Location']]).strip() if row[ix['Location']] is not None else ''
        city=str(row[ix['City_Name']]).strip() if row[ix['City_Name']] is not None else ''
        if not country or not city: continue
        raw_pop=row[ix['2025']]
        if raw_pop is None or str(raw_pop).strip() == '':
            blank += 1
            records.append({
                'city': city,
                'country': country,
                'wup_city_code': row[ix['City_Code']],
                'iso3': str(row[ix['ISO3_Code']]).strip() if row[ix['ISO3_Code']] is not None else '',
                'population': None,
                'population_status': 'unreported_subthreshold_in_F21',
                'source_row_index': source_row_index,
            })
        else:
            try: pop=float(raw_pop)
            except (TypeError,ValueError): continue
            if math.isfinite(pop) and pop >= 50: existing += 1
    if len(rows)-1 != EXPECTED_TOTAL: raise SystemExit(f'wup-f21-row-count:{len(rows)-1}:expected:{EXPECTED_TOTAL}')
    if existing != EXPECTED_EXISTING: raise SystemExit(f'wup-f21-existing-count:{existing}:expected:{EXPECTED_EXISTING}')
    if blank != EXPECTED_EXTENSION: raise SystemExit(f'wup-f21-extension-count:{blank}:expected:{EXPECTED_EXTENSION}')
    if len({(r['country'].casefold(),r['city'].casefold()) for r in records}) != EXPECTED_EXTENSION:
        raise SystemExit('wup-f21-extension-duplicate-identity')
    Path(output).write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'source':SOURCE_URL,'sha256':digest,'total_rows':EXPECTED_TOTAL,'existing_ge_50k':existing,'extension_rows':blank,'output':output},indent=2))

if __name__ == '__main__':
    import sys
    build(sys.argv[1],sys.argv[2])
